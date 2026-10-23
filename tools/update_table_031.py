# -*- coding: utf-8 -*-
"""按 0.31 解包更新《白卡植物标签表.xlsx》，并新增「按章节分类」工作表。

做四件事：
  1. 「白卡植物标签」：追加 0.31 新增的 14 株白卡（第10章 10 株 + 杂交乐园 4 株）；
  2. 「标签筛选」：同步追加这 14 株的逐标签行；
  3. 「图鉴说明」：从 Translate.zh.translation 抽出简介 / 属性（含获取方式）/ 图鉴故事；
  4. 新增工作表「按章节分类」：按「获取方式」把 202 株白卡分组（冒险第1~10章 / 杂交乐园 /
     在线关卡 / 商店购买 / 抽奖盒子 / 迷你游戏 / 挑战模式 / 我是僵尸 / 对战模式 / 限时活动 / 初始可用）。

数据全部来自解包（配置字段 + 组件集 + 图鉴文案），口径与老表一致（见 tools/validate_031.py 的比对）。

用法：python tools/update_table_031.py [--dry]
"""
import copy
import io
import json
import os
import re
import shutil
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
MOD = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import extract_031_facts as F          # noqa: E402
import zh_texts as Z                   # noqa: E402

XLSX = os.path.join(MOD, "data", "白卡植物标签表.xlsx")
XLSX_ROOT = os.path.join(os.path.dirname(os.path.dirname(MOD)), "白卡植物标签表.xlsx")

# ── 章节分组（与 runtime_src/WhiteCardChapterSources.cs 的 gen_chapter_sources.py 完全同口径）──
GROUPS = [
    ("A1", "第1章", 10), ("A2", "第2章", 20), ("A3", "第3章", 30), ("A4", "第4章", 40),
    ("A5", "第5章", 50), ("A6", "第6章", 60), ("A7", "第7章", 70), ("A8", "第8章", 80),
    ("A9", "第9章", 90), ("A10", "第10章", 100),
    ("HY", "杂交乐园", 200), ("ON", "在线关卡", 300), ("SH", "商店购买", 400),
    ("LT", "抽奖盒子", 500), ("MINI", "迷你游戏", 550), ("CHAL", "挑战模式", 560),
    ("IZM", "我是僵尸", 570), ("VS", "对战模式", 580), ("EVENT", "限时活动", 590),
    ("BA", "初始可用", 600),
]
GROUP_LABEL = {g: l for g, l, _ in GROUPS}
GROUP_WEIGHT = {g: w for g, _, w in GROUPS}


def classify(text):
    if not text:
        return "BA"
    m = re.match(r"通关冒险模式\s*(\d+)\s*[-－]\s*(\d+)", text)
    if m:
        ch = int(m.group(1))
        return "A%d" % ch if 1 <= ch <= 10 else "BA"
    if "杂交乐园" in text:
        return "HY"
    if text.startswith("在线关卡奖励"):
        return "ON"
    if text.startswith("商店购买"):
        return "SH"
    if "抽奖盒子" in text:
        return "LT"
    if "迷你游戏" in text or "砸罐子" in text or "幸运四叶草" in text or "坚果保龄球" in text:
        return "MINI"
    if "挑战模式" in text:
        return "CHAL"
    if "我是僵尸" in text:
        return "IZM"
    if "对战" in text:
        return "VS"
    if "每日挑战" in text or "活动" in text:
        return "EVENT"
    return "BA"


# ── 0.31 新增的 14 株：标签与「攻击/弹道」按《标签字典》判定依据 + 图鉴「特性/特点」逐株核定 ──
NEW_ROWS = {
    "PlantLilyPadMagic": dict(
        attack="射击", tags=["射击", "花盆", "莲叶", "魔法", "水生"],
        evidence="组件：FireComponent、SlotComponent；physiqueTypeFlags=4100（莲叶+魔法）；plantGridType=[3]（仅水面）"),
    "PlantIceKelp": dict(
        attack="—", tags=["近战", "缠绕", "冰冻", "冰", "水生"],
        evidence="组件：AttackComponent、TanglekelpComponent；plantGridType=[3]；图鉴「特性：水生，冰属性」+「消失时释放寒冰菇效果」⇒ 冰/冰冻"),
    "PlantChomperMagic": dict(
        attack="—", tags=["近战", "魔法"],
        evidence="组件：AttackComponent；physiqueTypeFlags=4096（魔法）；图鉴「不具备大嘴花的吞噬功能」⇒ 不打大嘴花标签"),
    "PlantTallnutSurf": dict(
        attack="—", tags=["坚果", "高坚果", "魔法", "冰", "水生"],
        evidence="physiqueTypeFlags=4097（坚果+魔法）；plantGridType=[3]；组件：CharacterTimerComponent（漩涡）；图鉴「特性：水生，冰属性，魔法」"),
    "PlantScaredyShroomElf": dict(
        attack="射击", tags=["射击", "胆小", "魔法", "夜行"],
        evidence="组件：FireComponent、ScaredComponent；physiqueTypeFlags=4096（魔法）；sleepTime=Day（夜行）"),
    "PlantSeaShroomMagic": dict(
        attack="射击", tags=["射击", "魔法", "夜行", "水生"],
        evidence="组件：FireComponent；physiqueTypeFlags=4096（魔法）；plantGridType=[3]；sleepTime=Day（夜行）"),
    "PlantFumeShroomMagic": dict(
        attack="—", tags=["近战", "喷菇", "魔法", "夜行"],
        evidence="组件：AttackComponent（喷菇系按近战口径）；physiqueTypeFlags=4096（魔法）；sleepTime=Day；名称含 FumeShroom⇒喷菇"),
    "PlantCattailMagic": dict(
        attack="猫尾", tags=["猫尾", "穿透", "魔法", "水陆"],
        evidence="弹道 SpikeMagic 的 fireMethodFlags=32（TRACK）；physiqueTypeFlags=4352（猫尾+魔法）；plantGridType=[2,3,6,5]（含水面且不止水面⇒水陆）"),
    "PlantPumpkinWizard": dict(
        attack="—", tags=["坚果", "魔法"],
        evidence="physiqueTypeFlags=4097（坚果+魔法）；组件：CharacterTimerComponent（每20秒投掷）；图鉴「不具备南瓜壳功能」⇒ 不打南瓜壳标签"),
    "PlantMagicStoneNut": dict(
        attack="—", tags=["坚果", "魔法"],
        evidence="physiqueTypeFlags=4097（坚果+魔法）；组件：HurtComponent 破损时触发全场魔法植物攻速加成"),
    "PlantPlanternCoffee": dict(
        attack="—", tags=["灯具", "咖啡豆"],
        evidence="physiqueTypeFlags=64（灯具）；组件：CharacterAabbAreaComponent（3×3 光环）；图鉴「特性：路灯，咖啡」"),
    "PlantMagnetPot": dict(
        attack="—", tags=["花盆", "磁力"],
        evidence="physiqueTypeFlags=2（花盆）、组件 SlotComponent；名称/玩法为磁力菇服务（无 MagnetComponent，磁力按名称判定）"),
    "PlantPotatoPumpkin": dict(
        attack="—", tags=["近战", "一次性", "成长", "南瓜壳", "坚果"],
        evidence="组件：AttackComponent、ExplodeComponent（一次性）、GrowUpComponent（成长）；physiqueTypeFlags=1（坚果）；图鉴「特性：南瓜壳」"),
    "PlantCoffeePot": dict(
        attack="—", tags=["花盆", "咖啡豆"],
        evidence="physiqueTypeFlags=10（咖啡豆+花盆）；组件：SlotComponent；图鉴「特性：底座，无土栽培，咖啡」"),
}


def style_like(ws, src_row, dst_row, ncols):
    for c in range(1, ncols + 1):
        s = ws.cell(row=src_row, column=c)
        d = ws.cell(row=dst_row, column=c)
        if s.has_style:
            d.font = copy.copy(s.font)
            d.border = copy.copy(s.border)
            d.fill = copy.copy(s.fill)
            d.alignment = copy.copy(s.alignment)
            d.number_format = s.number_format
            d.protection = copy.copy(s.protection)


def main():
    dry = "--dry" in sys.argv
    wb = openpyxl.load_workbook(XLSX)
    tr = F.load_translate()
    idx = F.packet_index()
    texts = Z.load()
    white = F.white_list()

    ws = wb["白卡植物标签"]
    existing = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[2]]
    have = {str(r[2]).strip() for r in existing}
    todo = [k for k in white if k not in have]
    print("现有 %d 株，解包 %d 株，待补 %d 株" % (len(existing), len(white), len(todo)))
    missing_meta = [k for k in todo if k not in NEW_ROWS]
    if missing_meta:
        print("!! 这些新卡没有手工核定标签：", missing_meta)
        return 1

    facts = {k: F.facts_of(k, idx, tr, "") for k in todo}
    seq = len(existing)
    last_ws_row = 1 + len(existing)

    # ── 1) 白卡植物标签 ──
    for k in todo:
        f = facts[k]
        m = NEW_ROWS[k]
        t = texts.get(f["cn"], {})
        seq += 1
        row = [
            seq, f["cn"], k,
            f.get("cost") or None,
            (round(float(f["cooldown"]), 2) if f.get("cooldown") else None),
            (int(float(f["hitpoints"])) if f.get("hitpoints") else 300),
            m["attack"],
            "、".join(m["tags"]),
            m["evidence"],
        ]
        r = ws.max_row + 1
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
        style_like(ws, last_ws_row, r, 9)
        print("  + [%d] %s %s  阳光=%s 冷却=%s 韧性=%s 攻/弹=%s 标签=%s"
              % (seq, f["cn"], k, row[3], row[4], row[5], m["attack"], row[7]))

    # ── 2) 标签筛选 ──
    wf = wb["标签筛选"]
    frow = wf.max_row
    for k in todo:
        f = facts[k]
        m = NEW_ROWS[k]
        for tag in m["tags"]:
            frow += 1
            vals = [seq_of(ws, k), f["cn"], tag,
                    f.get("cost") or None,
                    (round(float(f["cooldown"]), 2) if f.get("cooldown") else None),
                    (int(float(f["hitpoints"])) if f.get("hitpoints") else 300),
                    m["attack"]]
            for c, v in enumerate(vals, start=1):
                wf.cell(row=frow, column=c, value=v)
            style_like(wf, wf.max_row - 1 if False else frow - 1, frow, 7)

    # ── 3) 图鉴说明 ──
    wt = wb["图鉴说明"]
    trow = wt.max_row
    for k in todo:
        f = facts[k]
        t = texts.get(f["cn"], {})
        trow += 1
        vals = [seq_of(ws, k), f["cn"], k, t.get("intro", ""), t.get("detail", ""), t.get("story", "")]
        for c, v in enumerate(vals, start=1):
            wt.cell(row=trow, column=c, value=v)
        style_like(wt, trow - 1, trow, 6)

    # ── 4) 按章节分类（整表重建）──
    name = "按章节分类"
    if name in wb.sheetnames:
        del wb[name]
    wc = wb.create_sheet(name, wb.sheetnames.index("图鉴说明") + 1)
    head = ["章节分组", "排序权重", "序号", "植物名称", "植物键", "阳光", "冷却(秒)", "韧性",
            "攻击/弹道", "标签", "获取方式（图鉴原文）"]
    wc.append(head)
    for c in range(1, len(head) + 1):
        cell = wc.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")

    def acq_of(detail):
        if not detail:
            return ""
        m = re.search(r"获取方式[：:]\s*([^\n\r]*)", detail)
        return m.group(1).strip() if m else ""

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[2]:
            continue
        key = str(r[2]).strip()
        detail = None
        for tr_ in wt.iter_rows(min_row=2, values_only=True):
            if tr_ and tr_[2] and str(tr_[2]).strip() == key:
                detail = tr_[4]
                break
        acq = acq_of(detail)
        g = classify(acq)
        rows.append((GROUP_LABEL[g], GROUP_WEIGHT[g], r[0], r[1], key, r[3], r[4], r[5], r[6], r[7], acq))

    rows.sort(key=lambda x: (x[1], x[2] if isinstance(x[2], int) else 9999))
    for v in rows:
        wc.append(list(v))
    wc.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJK", [12, 10, 8, 16, 26, 8, 10, 8, 12, 40, 40]):
        wc.column_dimensions[col].width = w

    # ── 5) 数据说明 ──
    wd = wb["数据说明"]
    for r in wd.iter_rows(min_row=2, max_row=wd.max_row):
        k = str(r[0].value or "")
        if k == "数据来源":
            r[1].value = re.sub(r"\(0\.\d+\.\d+\)", "(0.31.0)", str(r[1].value))
    add = [
        ("白卡总数", "%d 株（White 卡池数组顺序；含 0.31 新增 %d 株）" % (len(white), len(todo))),
        ("本次更新", "0.31.0：补齐第10章 10 株（10-11~10-20）+ 杂交乐园 4 株；新增工作表「按章节分类」"),
        ("按章节分类", "按「图鉴说明」的「获取方式」分组：冒险第1~10章 / 杂交乐园 / 在线关卡 / 商店购买 / 抽奖盒子 / 迷你游戏 / 挑战模式 / 我是僵尸 / 对战模式 / 限时活动 / 初始可用"),
        ("文案抽取", "Translate.zh.translation（Godot OptimizedTranslation：RSCC + zstd 帧）→ 按 \\x00 字段边界取 [故事][中文名][简介][属性]"),
    ]
    have_k = {str(r[0].value or "") for r in wd.iter_rows(min_row=2)}
    for k, v in add:
        if k in have_k:
            for r in wd.iter_rows(min_row=2):
                if str(r[0].value or "") == k:
                    r[1].value = v
                    break
        else:
            wd.append([k, v])

    total = ws.max_row - 1
    print("\n表格统计：白卡 %d 株；标签筛选 %d 行；图鉴说明 %d 行；按章节分类 %d 行"
          % (total, wf.max_row - 1, wt.max_row - 1, wc.max_row - 1))
    from collections import Counter
    cnt = Counter(v[0] for v in rows)
    for g, label, _ in GROUPS:
        if cnt.get(label):
            print("   %-8s %3d 株" % (label, cnt[label]))

    if dry:
        print("\n--dry：不写盘")
        return 0
    wb.save(XLSX)
    shutil.copyfile(XLSX, XLSX_ROOT)
    print("\n已写入:", XLSX)
    print("已同步:", XLSX_ROOT)
    return 0


def seq_of(ws, key):
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[2] and str(r[2]).strip() == key:
            return r[0]
    return None


if __name__ == "__main__":
    raise SystemExit(main())
