# -*- coding: utf-8 -*-
"""校验：0.31 解包提取出的事实/文案 与 现有 xlsx 的 188 行逐项比对。

目的：在把新白卡写进表格之前，先证明提取口径与老表**完全一致**（否则新数据的可信度无从谈起）。

用法：python tools/validate_031.py
"""
import io
import json
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import extract_031_facts as F          # noqa: E402
import zh_texts as Z                   # noqa: E402

MOD = os.path.dirname(HERE)
XLSX = os.path.join(MOD, "data", "白卡植物标签表.xlsx")


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    rows = [r for r in wb["白卡植物标签"].iter_rows(min_row=2, values_only=True) if r and r[2]]
    texts = wb["图鉴说明"]
    zh = {r[2]: (r[3], r[4], r[5]) for r in texts.iter_rows(min_row=2, values_only=True) if r and r[2]}

    tr = F.load_translate()
    idx = F.packet_index()
    white = F.white_list()

    # 文案偏移口径验证：story 取「名字前一位」还是「属性块后一位」
    fields = Z._blob().split("\x00")
    pos = {}
    for i, f in enumerate(fields):
        if f and i + 2 < len(fields) and fields[i + 1] and fields[i + 2] \
                and ("获取方式" in fields[i + 2] or "韧性" in fields[i + 2]):
            pos.setdefault(f, i)

    stat = {"intro_hit": 0, "detail_hit": 0, "story_prev_hit": 0, "story_next_hit": 0,
            "n": 0, "cost_hit": 0, "cd_hit": 0, "hp_hit": 0, "name_hit": 0}
    diffs = []
    for r in rows:
        key = str(r[2]).strip()
        cn = str(r[1]).strip()
        f = F.facts_of(key, idx, tr, "")
        if not f:
            continue
        stat["n"] += 1
        if f["cn"] == cn:
            stat["name_hit"] += 1
        et_intro, et_detail, et_story = (zh.get(key) or ("", "", ""))
        i = pos.get(cn)
        if i is not None:
            intro = Z.strip_bb(fields[i + 1])
            detail = Z.strip_bb(fields[i + 2])
            prev = Z.strip_bb(fields[i - 1]) if i > 0 else ""
            nxt = Z.strip_bb(fields[i + 3]) if i + 3 < len(fields) else ""
            if intro == (et_intro or ""):
                stat["intro_hit"] += 1
            if detail == (et_detail or ""):
                stat["detail_hit"] += 1
            if prev == (et_story or ""):
                stat["story_prev_hit"] += 1
            if nxt == (et_story or ""):
                stat["story_next_hit"] += 1
            if intro != (et_intro or "") or detail != (et_detail or ""):
                diffs.append((key, "文案", (et_intro or "")[:40], intro[:40]))
        # 数值
        cost = str(f.get("cost") or "").strip()
        cd = str(f.get("cooldown") or "").strip()
        hp = str(f.get("hitpoints") or "").strip()
        if cost and cost == str(r[3]).strip():
            stat["cost_hit"] += 1
        if cd and abs(float(cd) - float(str(r[4]).strip() or 0)) < 1e-6:
            stat["cd_hit"] += 1
        if hp and abs(float(hp) - float(str(r[5]).strip() or 0)) < 1e-6:
            stat["hp_hit"] += 1

    n = stat["n"]
    print("比对行数:", n)
    for k in ["name_hit", "intro_hit", "detail_hit", "story_prev_hit", "story_next_hit",
              "cost_hit", "cd_hit", "hp_hit"]:
        print("  %-16s %3d / %d" % (k, stat[k], n))
    print("\n前 10 处文案差异：")
    for d in diffs[:10]:
        print("   ", d[0], "\n      表:", d[2], "\n      解:", d[3])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
